"""Reading XLSX/CSV exports into raw rows.

Values are returned exactly as the file provides them — no type coercion, no
NaN semantics, no silent cleanup. All interpretation happens later, in the
normalization layer, so the raw source values stay reconstructable.
"""

from __future__ import annotations

import csv
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

# Export files sometimes carry a title or a date line above the real header.
MAX_HEADER_SCAN_ROWS = 20
MIN_HEADER_CELLS = 2

CSV_ENCODINGS = ("utf-8-sig", "cp1251")
CSV_DELIMITERS = (";", ",", "\t")

XLSX_SUFFIXES = frozenset({".xlsx", ".xlsm"})
CSV_SUFFIXES = frozenset({".csv", ".tsv", ".txt"})


class SourceFileError(Exception):
    """Raised when a file cannot be read as a table at all."""


@dataclass(frozen=True)
class SourceRow:
    """One data row with its 1-based line number in the source file."""

    number: int
    values: list[Any]


@dataclass(frozen=True)
class SourceTable:
    """Headers plus data rows of a single sheet or CSV file."""

    headers: list[str]
    rows: list[SourceRow]
    header_row_number: int
    sheet_name: str | None = None
    encoding: str | None = None
    delimiter: str | None = None


def read_table(path: Path, *, sheet: str | None = None) -> SourceTable:
    """Read an XLSX or CSV file into a :class:`SourceTable`."""
    if not path.exists():
        raise SourceFileError(f"file not found: {path}")
    suffix = path.suffix.lower()
    if suffix in XLSX_SUFFIXES:
        return _read_xlsx(path, sheet=sheet)
    if suffix in CSV_SUFFIXES:
        return _read_csv(path)
    raise SourceFileError(
        f"unsupported file type '{suffix}'; expected one of "
        f"{sorted(XLSX_SUFFIXES | CSV_SUFFIXES)}"
    )


def _is_empty(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _find_header_index(rows: list[list[Any]]) -> int:
    """Index of the first row that looks like a header row."""
    for index, row in enumerate(rows[:MAX_HEADER_SCAN_ROWS]):
        if sum(not _is_empty(cell) for cell in row) >= MIN_HEADER_CELLS:
            return index
    raise SourceFileError(
        f"no header row found in the first {MAX_HEADER_SCAN_ROWS} rows "
        f"(a header needs at least {MIN_HEADER_CELLS} non-empty cells)"
    )


def _build_table(
    raw_rows: list[list[Any]],
    *,
    sheet_name: str | None = None,
    encoding: str | None = None,
    delimiter: str | None = None,
) -> SourceTable:
    if not raw_rows:
        raise SourceFileError("file contains no rows")

    header_index = _find_header_index(raw_rows)
    headers = [("" if _is_empty(cell) else str(cell).strip()) for cell in raw_rows[header_index]]

    rows = [
        SourceRow(number=header_index + offset + 2, values=values)
        for offset, values in enumerate(raw_rows[header_index + 1 :])
        if any(not _is_empty(cell) for cell in values)
    ]
    return SourceTable(
        headers=headers,
        rows=rows,
        header_row_number=header_index + 1,
        sheet_name=sheet_name,
        encoding=encoding,
        delimiter=delimiter,
    )


def _read_xlsx(path: Path, *, sheet: str | None) -> SourceTable:
    workbook = load_workbook(filename=path, read_only=True, data_only=True)
    try:
        if sheet is not None:
            if sheet not in workbook.sheetnames:
                raise SourceFileError(
                    f"sheet {sheet!r} not found; available: {workbook.sheetnames}"
                )
            worksheet = workbook[sheet]
        else:
            worksheet = workbook[workbook.sheetnames[0]]
        raw_rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
    finally:
        workbook.close()
    return _build_table(raw_rows, sheet_name=worksheet.title)


def _detect_delimiter(sample: str) -> str:
    try:
        return csv.Sniffer().sniff(sample, delimiters="".join(CSV_DELIMITERS)).delimiter
    except csv.Error:
        first_line = sample.splitlines()[0] if sample.splitlines() else ""
        return max(CSV_DELIMITERS, key=first_line.count)


def _read_csv(path: Path) -> SourceTable:
    text: str | None = None
    used_encoding: str | None = None
    for encoding in CSV_ENCODINGS:
        try:
            text = path.read_text(encoding=encoding)
            used_encoding = encoding
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise SourceFileError(f"could not decode {path.name} using {list(CSV_ENCODINGS)}")

    delimiter = _detect_delimiter(text[:8192])
    raw_rows = [list(row) for row in csv.reader(text.splitlines(), delimiter=delimiter)]
    return _build_table(raw_rows, encoding=used_encoding, delimiter=delimiter)
